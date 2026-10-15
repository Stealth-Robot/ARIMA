"""Convert the Lettuce Simuls xlsx export into app/data/simul_seed.json.

Usage: python scripts/simul_xlsx_to_seed.py <path/to/Lettuce Simuls.xlsx>

Does not touch the database; the startup migration imports the JSON.
"""
import json
import os
import re
import sys

import openpyxl

SEED_PATH = os.path.join(os.path.dirname(__file__), '..', 'app', 'data', 'simul_seed.json')

STATUS_TYPES = [
    {'key': 'watched', 'label': 'Watched', 'color_hex': '#4a7c59', 'applies_to': 'watch', 'sort_order': 1},
    {'key': 'interested', 'label': 'Interested', 'color_hex': '#4a6f9c', 'applies_to': 'planning', 'sort_order': 2},
    {'key': 'not_interested', 'label': 'Not interested', 'color_hex': '#8a4a4a', 'applies_to': 'planning', 'sort_order': 3},
    {'key': 'other', 'label': 'Other', 'color_hex': '#6f5d94', 'applies_to': 'both', 'sort_order': 4},
]

# Planning title fills, per the sheet's own legend in L1:M2.
CATEGORY_BY_FILL = {
    'FFD9D2E9': 'tv',
    'FFFCE5CD': 'movie',
    'FFB6D7A8': 'upcoming_seasonal',
    'theme9': 'upcoming_movie',
}
# Watch sheets reuse only the TV/movie fills; other title highlights there carry no category.
WATCH_CATEGORY_BY_FILL = {'FFD9D2E9': 'tv', 'FFFCE5CD': 'movie'}
NO_STATUS_FILLS = {None, 'FFFFFFFF'}
PLANNING_NOT_INTERESTED_FILL = 'FFEA9999'
WATCHED_FILLS = {'FFFF0000', 'FFE06666'}
WATCH_SHEET_RE = re.compile(r'^(OngoingPaused|Completed|Dropped) (\d{4})$')
BUCKET_BY_PREFIX = {'OngoingPaused': 'ongoing', 'Completed': 'completed', 'Dropped': 'dropped'}
HISTORICAL_MEMBERS = {'Amy', 'Syd'}


def fill_of(cell):
    if not (cell.fill and cell.fill.fill_type):
        return None
    fg = cell.fill.fgColor
    if fg.type == 'rgb':
        return fg.rgb
    if fg.type == 'theme':
        return 'theme%s' % fg.theme
    return None


def text(value):
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def comment_of(cell):
    return text(cell.comment.text) if cell.comment else None


def status_for(cell, planning):
    """Return (status_key, note, youtube_url) for one participant cell, or None when blank."""
    fill = fill_of(cell)
    value = text(cell.value)
    note_parts = []
    youtube = None
    is_not_interested = False
    if value:
        if 'youtube.com' in value or 'youtu.be' in value:
            youtube = value
        elif value.lower() == 'not interested':
            is_not_interested = True
        else:
            note_parts.append(value)
    if comment_of(cell):
        note_parts.append(comment_of(cell))
    note = '\n'.join(note_parts) or None

    if planning:
        if fill == PLANNING_NOT_INTERESTED_FILL or is_not_interested:
            key = 'not_interested'
        elif fill not in NO_STATUS_FILLS or note:
            key = 'other'
        else:
            key = None
    else:
        if is_not_interested:
            key = 'not_interested'
        elif fill in WATCHED_FILLS:
            key = 'watched'
        elif fill not in NO_STATUS_FILLS or note:
            key = 'other'
        else:
            key = None
    if key is None:
        return None if youtube is None else (None, None, youtube)
    return key, note, youtube


def convert(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path)
    shows = []
    for ws in wb.worksheets:
        if ws.title == 'Simul Planning':
            planning, bucket, year, header_row = True, 'planning', None, 2
        else:
            m = WATCH_SHEET_RE.match(ws.title)
            if not m:
                print('Skipping unrecognised sheet %r' % ws.title, file=sys.stderr)
                continue
            planning, bucket, year, header_row = False, BUCKET_BY_PREFIX[m.group(1)], int(m.group(2)), 1

        members = {}
        for c in ws[header_row][1:]:
            name = text(c.value)
            if not name or name.isupper():
                break  # planning's category legend starts after the member columns
            members[c.column] = name
        extra_col = max(members) + 1

        for row in ws.iter_rows(min_row=header_row + 1):
            title = text(row[0].value)
            if not title:
                continue
            show = {
                'title': title,
                'category': (CATEGORY_BY_FILL if planning else WATCH_CATEGORY_BY_FILL).get(
                    fill_of(row[0]), 'uncategorized'),
                'bucket': bucket,
                'year': year,
                'youtube_url': None,
                'url': text(row[0].hyperlink.target) if row[0].hyperlink else None,
                'notes': comment_of(row[0]),
                'years_running': None if planning else text(row[extra_col - 1].value) if len(row) >= extra_col else None,
                'sort_order': len(shows) + 1,
                'statuses': [],
            }
            for cell in row[1:]:
                name = members.get(cell.column)
                if not name:
                    continue
                result = status_for(cell, planning)
                if result is None:
                    continue
                key, note, youtube = result
                if youtube:
                    show['youtube_url'] = youtube
                if key is None:
                    continue
                historical = name in HISTORICAL_MEMBERS
                show['statuses'].append({
                    'user': None if historical else name,
                    'member_name': name if historical else None,
                    'status_key': key,
                    'note': note,
                })
            shows.append(show)
    return {'status_types': STATUS_TYPES, 'shows': shows}


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    data = convert(sys.argv[1])
    with open(SEED_PATH, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    print('Wrote %d shows to %s' % (len(data['shows']), os.path.normpath(SEED_PATH)))


if __name__ == '__main__':
    main()
